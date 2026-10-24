import io
import tempfile
import unittest
from datetime import date
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from pipeline.participants import parse_form4, annotate, sync, read_snapshot, workbook_paths
from pipeline.participants_ui import export_participants
from pipeline.repository import LocalRepository


def workbook(sheets):
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        for name,data in sheets.items():
            data.to_excel(writer,sheet_name=name,index=False)
    return out.getvalue()


def form4_data(status='OK',buys=0,sells=0,unknown=0):
    return {'Ticker Summary':pd.DataFrame([{'Ticker':'META','Analysis Start':'2026-09-25','Analysis End':'2026-10-24','Buy Count':buys,'Sell Count':sells}]),
            'Form4 Detail':pd.DataFrame(columns=['Ticker','Transaction Code','Trade Date','Insider']),
            'Run Log':pd.DataFrame([{'Ticker':'META','Status':status,'Issuer Unknown Filings':unknown}])}


class Source(LocalRepository):
    def revision(self):
        return 'source-sha'


class ParticipantTests(unittest.TestCase):
    def test_form4_matches_source_npr_and_quality(self):
        sheets=parse_form4(workbook(form4_data(buys=2,sells=7)))
        row=sheets['Ticker Summary'].iloc[0]
        self.assertAlmostEqual(row['NPR'],-5/9)
        self.assertEqual(row['內部人交易動向'],'近期淨賣出')
        self.assertEqual(row['資料品質'],'OK')
        with self.assertRaises(ValueError):
            parse_form4(workbook(form4_data(buys=-1)))

    def test_failed_or_unknown_issuer_does_not_claim_no_activity(self):
        for status,unknown in [('ERROR',0),('OK',1)]:
            summary=parse_form4(workbook(form4_data(status=status,unknown=unknown)))['Ticker Summary']
            news=pd.DataFrame([{'Market':'US','Ticker':'META','Level':'5','news_id':'a'}])
            result=annotate(news,({'datasets':{}},{'form4':summary}),date(2026,10,24))
            self.assertIn('資料待覆核',result.iloc[0]['Form4 內部人動向'])
            self.assertNotIn('無P/S交易',result.iloc[0]['Form4 內部人動向'])
            self.assertEqual(result.iloc[0]['Level'],'5')

    def test_missing_uncovered_unknown_market_and_no_activity_are_distinct(self):
        summary=parse_form4(workbook(form4_data()))['Ticker Summary']
        rows=[{'Market':'US','Ticker':'META'},{'Market':'US','Ticker':'IBM'},{'Market':'US','Ticker':''},{'Market':'HK','Ticker':'META'}]
        result=annotate(pd.DataFrame(rows),({'datasets':{}},{'form4':summary}),date(2026,10,24))
        self.assertEqual(len(result),4)
        self.assertIn('無P/S交易',result.iloc[0]['Form4 內部人動向'])
        self.assertEqual(result.iloc[1]['Form4 內部人動向'],'不在本次追蹤清單')
        self.assertEqual(result.iloc[2]['Form4 內部人動向'],'股票代號待核對')
        self.assertEqual(result.iloc[3]['Form4 內部人動向'],'非美股／不適用')
        self.assertEqual(result.iloc[0]['13F 機構動向'],'尚無已發布成果')

    def test_sync_keeps_last_good_when_new_workbook_invalid(self):
        with tempfile.TemporaryDirectory() as folder, tempfile.TemporaryDirectory() as target_dir:
            source=Source(folder);target=LocalRepository(target_dir)
            path='insider_form4/2026-10-24/Form4_2026-10-24_30D_Full_Universe.xlsx'
            p=Path(folder)/path;p.parent.mkdir(parents=True);p.write_bytes(workbook(form4_data(buys=1)))
            first=sync(target,source)
            self.assertEqual(first['datasets']['form4']['status'],'ready')
            self.assertEqual(first['datasets']['13f']['status'],'not_published')
            p.write_bytes(workbook(form4_data(status='ERROR')))
            second=sync(target,source)
            self.assertEqual(second['datasets']['form4']['source_status'],'unavailable')
            self.assertEqual(read_snapshot(target)[1]['form4'].iloc[0]['Buy Count'],'1')

    def test_export_appends_signals_and_form4_details_without_row_expansion(self):
        sheets=parse_form4(workbook(form4_data(buys=2,sells=7)))
        snapshot=({'datasets':{'form4':{'source_sha':'abc','source_status':'available'}}},{'form4':sheets['Ticker Summary']})
        news=pd.DataFrame([{'Market':'US','Ticker':'META','Title':'Meta test','日期':'2026-10-24','Level':'3','news_id':str(i)} for i in range(2)])
        out=export_participants(news,snapshot,{'form4':{'sheets':sheets}})
        book=load_workbook(io.BytesIO(out))
        self.assertEqual(book['完整新聞'].max_row,3)
        self.assertEqual([c.value for c in book['完整新聞'][1]][-2:],['13F 機構動向','Form4 內部人動向'])
        self.assertIn('Form4 Form4 Detail',book.sheetnames)
        self.assertTrue(book['完整新聞']['A2'].fill.fgColor.rgb.endswith('FFF9C4'))

    def test_paths_use_source_folder_and_filename_dates(self):
        paths=['institutional_13f/2026Q2/13F_2026Q2_Full_Universe.xlsx','insider_form4/2026-10-24/Form4_2026-10-24_30D_Full_Universe.xlsx','insider_form4/2026-10-24/Form4_2026-10-23_30D_Full_Universe.xlsx']
        self.assertEqual(len(workbook_paths(paths,'13f')),1)
        self.assertEqual(len(workbook_paths(paths,'form4')),1)

    def test_manual_ticker_edit_refreshes_issuer_signals_in_saved_excel(self):
        from pipeline.records import normalize, read_days, manual_save
        from pipeline.repository import json_bytes
        with tempfile.TemporaryDirectory() as folder:
            repo=LocalRepository(folder)
            summary=parse_form4(workbook(form4_data(buys=2,sells=7)))['Ticker Summary']
            ibm=summary.iloc[0].copy();ibm['Ticker']='IBM';ibm['Buy Count']=3;ibm['Sell Count']=0
            summary=pd.concat([summary,pd.DataFrame([ibm])],ignore_index=True)
            news=normalize(pd.DataFrame([{'Market':'US','Ticker':'META','Company':'Meta Platforms','Title':'Meta and IBM credit risk','Published Time':'2026-10-24T10:00:00+08:00','Level':'3','URL':'https://example.com/edit'}]))
            files={'news_daily/2026-10-24.csv':news.to_csv(index=False).encode(),
                   'institutional_insider/latest.json':json_bytes({'datasets':{'form4':{'summary_path':'institutional_insider/form4/summary.csv','source_status':'available'}}}),
                   'institutional_insider/form4/summary.csv':summary.to_csv(index=False).encode()}
            repo.transaction(lambda read,rev:(files,None),'fixture')
            before=read_days(repo,date(2026,10,24),date(2026,10,24))
            edited=before.copy();edited['Ticker']='IBM'
            self.assertEqual(manual_save(repo,before,edited),1)
            saved=pd.read_csv(Path(folder)/'news_daily/2026-10-24.csv').iloc[0]
            self.assertIn('淨買入',saved['Form4 內部人動向'])
            book=load_workbook(io.BytesIO(repo.read('news_daily/2026-10-24.xlsx')))
            self.assertIn('淨買入',book['完整新聞'].cell(2,book['完整新聞'].max_column).value)


if __name__=='__main__':
    unittest.main()
