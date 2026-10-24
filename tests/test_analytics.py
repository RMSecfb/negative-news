import io
import unittest
import pandas as pd
from openpyxl import load_workbook
from pipeline.analytics import frequency_tables, join_sec, parse_sec, combined_excel, sec_paths


class AnalyticsTests(unittest.TestCase):
    def test_frequency_uses_observed_days_and_market_issuer_identity(self):
        rows = [{'日期': '2026-10-20', 'Market': 'US', 'Ticker': 'META', 'Level': '3', 'news_id': 'old'}]
        rows += [{'日期': '2026-10-23', 'Market': 'US', 'Ticker': 'META', 'Level': '3', 'news_id': str(i)} for i in range(2)]
        rows += [dict(rows[-1])]  # Duplicate issuer article is one item.
        rows += [{'日期': '2026-10-23', 'Market': 'US', 'Ticker': 'GETY', 'Level': '', 'news_id': 'pending'}]
        counts, spikes, days = frequency_tables(pd.DataFrame(rows))
        self.assertEqual(days, ['2026-10-20', '2026-10-23'])
        self.assertEqual(counts['則數'].sum(), 3)
        self.assertEqual(spikes.iloc[0]['過去平均每日'], 1)
        self.assertEqual(spikes.iloc[0]['Ticker'], 'META')

    def test_sec_join_no_cross_market_match_or_fanout_and_export(self):
        summary = pd.DataFrame([{'Ticker':'META', 'Quarter':'2026Q2', 'Quarter End':'2026-06-30', 'Company':'Meta'}])
        news = pd.DataFrame([{'Market':market,'Ticker':'META','Title':'Meta news', 'Level':'3', '日期':'2026-10-23', 'news_id':str(i)} for i,market in enumerate(['US','HK','US'])])
        sheets = {'Ticker Summary':summary,'13F Detail':pd.DataFrame([{'Institution':'Example', 'Ticker':'META'}]),'Run Log':pd.DataFrame([{'Status':'OK'}])}
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer,engine='openpyxl') as writer:
            for name, df in sheets.items():
                df.to_excel(writer,sheet_name=name,index=False)
        parsed = parse_sec(buffer.getvalue())
        joined = join_sec(news,parsed['Ticker Summary'])
        self.assertEqual(len(joined),3)
        self.assertTrue(pd.isna(joined.iloc[1]['13F · Quarter']))
        wb = load_workbook(io.BytesIO(combined_excel(news,parsed,{'SHA':'test'})))
        self.assertIn('完整新聞', wb.sheetnames)
        self.assertIn('13F季度摘要', wb.sheetnames)
        self.assertEqual(wb['新聞與13F對照'].max_row,4)

    def test_sec_paths_and_duplicate_rejected(self):
        self.assertEqual(sec_paths(['institutional_13f/2026Q2/13F_2026Q2_Full_Universe.xlsx','other.xlsx']),['institutional_13f/2026Q2/13F_2026Q2_Full_Universe.xlsx'])
        summary = pd.DataFrame([{'Ticker':'META','Quarter':'2026Q2','Quarter End':'2026-06-30'}]*2)
        buffer=io.BytesIO()
        with pd.ExcelWriter(buffer,engine='openpyxl') as writer:
            summary.to_excel(writer,sheet_name='Ticker Summary',index=False)
            pd.DataFrame().to_excel(writer,sheet_name='13F Detail')
            pd.DataFrame().to_excel(writer,sheet_name='Run Log')
        with self.assertRaises(ValueError):
            parse_sec(buffer.getvalue())
