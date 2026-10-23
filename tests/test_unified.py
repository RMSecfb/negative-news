import io
import tempfile
import unittest
from pathlib import Path
from datetime import datetime, date, timedelta
from unittest.mock import Mock, patch
import pandas as pd
from openpyxl import load_workbook
from pipeline.integrated import ROOT, execute, daily_request, apply_answer, collect
from pipeline.issuers import IssuerIndex, company_frame, load_companies
from pipeline.records import TPE, normalize, merge_news, read_days, manual_save, excel_bytes
from pipeline.repository import LocalRepository, read_json
from pipeline.events import load_rules


def item(title='Credit default swap spreads surge', ticker='META', day='2026-10-23', url='https://example.com/story?utm_source=x'):
    return {'Ticker': ticker, 'Company': 'Meta Platforms' if ticker else '', 'Market': 'US', 'Published Time': day + 'T10:00:00+08:00', 'Title': title, 'URL': url, 'Source': 'Test', '來源方法': '方法一'}


def answer(record, codes, key, labels):
    return {'company_mentioned': 'Meta', 'event_codes': ['CDS_SPREAD_WIDENING'], 'event_status': 'MARKET_IMPLIED', 'confirmed_default': False, 'evidence': record['Title'], 'summary_zh': '信用風險升高', 'relevant': True, 'requires_review': False}, {'input_tokens': 100, 'output_tokens': 40}, ''


class UnifiedTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.repo = LocalRepository(self.tmp.name)
        self.request = {'id': 'test-first', 'start': '2026-10-23T00:00:00+08:00', 'end': '2026-10-24T00:00:00+08:00', 'market': 'US', 'scope': 'all', 'lists': [], 'methods': ['方法一', '方法二'], 'ai': True, 'max_llm': 2}

    def tearDown(self):
        self.tmp.cleanup()

    def test_multi_list_and_hk_zero(self):
        companies = load_companies(ROOT, ['DowJones30', 'SP500'])
        self.assertEqual(len(companies), len(companies.drop_duplicates(['Market', 'Ticker'])))
        self.assertGreaterEqual(len(companies), 500)
        hk = company_frame(pd.DataFrame([{'Ticker': '00700', 'Company': 'Example'}]), 'HK')
        self.assertEqual(hk.iloc[0]['Ticker'], '0700.HK')

    def test_meta_resolves_and_ambiguous_does_not(self):
        index = IssuerIndex(load_companies(ROOT, ['SP500']))
        match, _ = index.resolve('Meta', 'US')
        self.assertEqual(match['Ticker'], 'META')
        conflict = IssuerIndex(pd.DataFrame([{'Market':'US','Ticker':'A','Company':'Same Inc'}, {'Market':'US','Ticker':'B','Company':'Same Corp'}]))
        self.assertIsNone(conflict.resolve('Same', 'US')[0])
        self.assertIsNone(index.resolve('Meta', 'HK')[0])

    def test_duplicate_methods_union_and_related_reports_retained(self):
        a = pd.DataFrame([item()])
        b = pd.DataFrame([{**item(url='https://example.com/story'), '來源方法':'方法二'}])
        out, stats = merge_news(a, b)
        self.assertEqual(len(out), 1)
        self.assertEqual(out.iloc[0]['來源方法'], '方法一|方法二')
        newer = item(title='Another report about CDS', url='https://other.com/story')
        out, _ = merge_news(out, pd.DataFrame([newer]))
        self.assertEqual(len(out), 2)

    def test_two_issuers_same_article_stay_separate(self):
        a = pd.DataFrame([item()])
        b = pd.DataFrame([{**item(ticker='AAPL'), 'Company':'Apple'}])
        out, _ = merge_news(a, b)
        self.assertEqual(len(out), 2)

    def test_ai_on_cached_and_meta_enrichment(self):
        collector = lambda *args: (pd.DataFrame([item(ticker='')]), [], ['mock collectors'])
        analyzer = Mock(side_effect=answer)
        first = execute(self.repo, self.request, api_key='test', collector=collector, analyzer=analyzer, interval=0)
        self.assertEqual(first['llm_calls'], 1)
        frame = read_days(self.repo, date(2026,10,23), date(2026,10,23))
        self.assertEqual(frame.iloc[0]['Ticker'], 'META')
        second = execute(self.repo, {**self.request,'id':'test-second'}, api_key='test', collector=collector, analyzer=analyzer, interval=0)
        self.assertEqual(second['llm_calls'], 0)
        self.assertEqual(analyzer.call_count, 1)
        self.assertEqual(len(read_days(self.repo,date(2026,10,23),date(2026,10,23))),1)

    def test_quota_queues_without_rule_becoming_final(self):
        data = pd.DataFrame([item(),item(title='Second debt story',url='https://example.com/second')])
        report = execute(self.repo, {**self.request,'max_llm':1}, api_key='test', collector=lambda *a:(data,[],[]), analyzer=answer, interval=0)
        self.assertEqual(report['llm_calls'],1)
        self.assertEqual(report['pending'],1)
        result=read_days(self.repo,date(2026,10,23),date(2026,10,23))
        queued=result[result['分類狀態']=='待補分析'].iloc[0]
        self.assertEqual(queued['Level'],'')
        self.assertEqual(queued['需要人工覆核'],'False')

    def test_ai_off_never_calls_model(self):
        analyzer=Mock()
        execute(self.repo,{**self.request,'ai':False},collector=lambda *a:(pd.DataFrame([item()]),[],[]),analyzer=analyzer,interval=0)
        analyzer.assert_not_called()
        frame=read_days(self.repo,date(2026,10,23),date(2026,10,23))
        self.assertEqual(frame.iloc[0]['最終判斷來源'],'規則')
        self.assertIn('Rule_Event_Code',frame)

    def test_period_query_and_excel_style(self):
        for day, level in [('2026-10-22','3'),('2026-10-23','4'),('2026-10-24','5')]:
            frame=normalize(pd.DataFrame([{**item(day=day,url='https://example.com/'+day),'Level':level}]))
            target=Path(self.tmp.name)/'news_daily'/f'{day}.csv'
            target.parent.mkdir(exist_ok=True)
            frame.to_csv(target,index=False)
        found=read_days(self.repo,date(2026,10,22),date(2026,10,24))
        self.assertEqual(len(found),3)
        workbook=load_workbook(io.BytesIO(excel_bytes(found)))
        self.assertIn('待補分析',workbook.sheetnames)
        colors={workbook['完整新聞'].cell(i,1).fill.fgColor.rgb[-6:] for i in (2,3,4)}
        self.assertEqual(colors,{'FFF9C4','FFE0B2','FFCDD2'})

    def test_manual_review_survives_rerun_and_conflict(self):
        collector=lambda *a:(pd.DataFrame([item()]),[],[])
        execute(self.repo,self.request,api_key='test',collector=collector,analyzer=answer,interval=0)
        original=read_days(self.repo,date(2026,10,23),date(2026,10,23))
        edited=original.copy()
        edited.loc[0,'Level']='4'
        edited.loc[0,'覆核註記']='confirmed by reviewer'
        self.assertEqual(manual_save(self.repo,original,edited,'test'),1)
        execute(self.repo,{**self.request,'id':'rerun'},api_key='test',collector=collector,analyzer=answer,interval=0)
        final=read_days(self.repo,date(2026,10,23),date(2026,10,23))
        self.assertEqual(final.iloc[0]['Level'],'4')
        self.assertEqual(final.iloc[0]['最終判斷來源'],'人工覆核')
        with self.assertRaisesRegex(ValueError,'新的人工覆核'):
            manual_save(self.repo,original,edited)

    def test_confirmed_default_guard(self):
        rules=load_rules(ROOT/'Parameter_Event.xlsx')
        index=IssuerIndex(load_companies(ROOT,['SP500']))
        a=answer(item(),[],'',{})[0]
        a.update(event_codes=['DEFAULT'],confirmed_default=False,event_status='RUMORED')
        out=apply_answer(item(title='Company might default'),a,rules,index)
        self.assertEqual(out['Event_Code'],'')
        self.assertEqual(out['分類狀態'],'待人工覆核')

    def test_daily_full_sp500_window(self):
        r=daily_request(datetime(2026,10,23,8,21,tzinfo=TPE))
        self.assertEqual(r['lists'],['SP500'])
        self.assertEqual(r['methods'],['方法一','方法二'])
        self.assertTrue(r['ai'] and r['discover_market'])
        self.assertEqual(datetime.fromisoformat(r['end'])-datetime.fromisoformat(r['start']),timedelta(days=1))

    def test_schedule_classifies_discovery_outside_sp500(self):
        request = daily_request(datetime(2026,10,24,8,21,tzinfo=TPE))
        request['id'] = 'market-outside-index'
        outside = {**item(ticker=''), 'Title': 'Outside issuer faces debt risk', 'Company': 'Outside issuer'}
        analyzer = Mock(side_effect=answer)
        result = execute(self.repo, request, api_key='test', collector=lambda *a:(pd.DataFrame([outside]),[],[]), analyzer=analyzer, interval=0)
        self.assertEqual(result['llm_calls'], 1)
        analyzer.assert_called_once()

    def test_export_log_survives_manual_save(self):
        execute(self.repo,self.request,api_key='test',collector=lambda *a:(pd.DataFrame([item()]),[],[]),analyzer=answer,interval=0)
        original=read_days(self.repo,date(2026,10,23),date(2026,10,23))
        edited=original.copy()
        edited.loc[0,'Level']='4'
        manual_save(self.repo,original,edited)
        workbook=load_workbook(io.BytesIO(self.repo.read('news_daily/2026-10-23.xlsx')))
        self.assertEqual(workbook['執行紀錄'].max_row,2)

    def test_collect_full_list_and_deduplicates_before_scoring(self):
        import legacy_engine as engine
        selected=load_companies(ROOT,['SP500'])
        data=pd.DataFrame([item()])
        req={**self.request,'scope':'list','lists':['SP500']}
        def score(frame):
            self.assertEqual(len(frame),1)
            out=frame.copy()
            out['Title']='Translated title'
            out['FinBERT']='-0.5'
            return out,pd.DataFrame()
        with patch.object(engine,'initialize'), patch.object(engine,'fetch_method1_news',return_value=(data.copy(),[],{})) as first, patch.object(engine,'fetch_method2_raw',return_value=(data.copy(),[],1)) as second, patch.object(engine,'score_method2_finbert',side_effect=score), patch('pipeline.integrated.event_discovery') as discovery:
            result,errors,stages=collect(ROOT,req,selected)
        self.assertGreaterEqual(len(first.call_args.args[0]),500)
        self.assertEqual(len(second.call_args.args[0]),len(selected))
        self.assertEqual(result.iloc[0]['Title'],item()['Title'])
        self.assertEqual(result.iloc[0]['FinBERT_Title'],'Translated title')
        self.assertEqual(result.iloc[0]['來源方法'],'方法一|方法二')
        discovery.assert_not_called()

    def test_symbol_punctuation_does_not_collide(self):
        companies=pd.DataFrame([{'Market':'US','Ticker':'META','Company':'Meta Platforms'}, {'Market':'US','Ticker':'MET$A','Company':'MetLife Preferred'}])
        index=IssuerIndex(companies)
        self.assertEqual(index.resolve('Meta','US')[0]['Ticker'],'META')
        self.assertEqual(index.resolve('MET$A','US')[0]['Ticker'],'MET$A')

    def test_cache_is_issuer_scoped(self):
        analyzer=Mock(side_effect=answer)
        execute(self.repo,self.request,api_key='test',collector=lambda *a:(pd.DataFrame([item()]),[],[]),analyzer=analyzer,interval=0)
        other={**item(ticker='AAPL'),'Company':'Apple'}
        result=execute(self.repo,{**self.request,'id':'second-issuer'},api_key='test',collector=lambda *a:(pd.DataFrame([other]),[],[]),analyzer=analyzer,interval=0)
        self.assertEqual(result['llm_calls'],1)
        self.assertEqual(analyzer.call_count,2)


if __name__=='__main__':
    unittest.main()
