"""Read existing Parameter_Event.xlsx; keyword hits are candidates, not final decisions."""
from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook

SYSTEM_CODES={'UNKNOWN','NOT_RELEVANT','NON_CORE_SOURCE'}
CDS_PAT = re.compile(r'\b(?:credit\s+default\s+swaps?|cds(?:\s+spreads?)?)\b|信用違約交換|信用違約互換|信用違約掉期', re.I)
ACTUAL_DEFAULT = re.compile(r'\b(?:(?:defaults?|defaulted)\s+on\s+(?:its\s+)?(?:loans?|debts?|bonds?|payments?)|missed?\s+(?:an?\s+)?(?:interest|debt|bond)\s+payment|failed?\s+to\s+repay|payment\s+default|event\s+of\s+default)\b|實際違約|未能償還|未支付利息|債務違約', re.I)


def load_rules(path: str|Path):
    wb=load_workbook(path,read_only=True,data_only=True)
    if 'Event' not in wb.sheetnames: raise ValueError('Parameter_Event.xlsx 缺少 Event 工作表')
    ws=wb['Event']; headers=[str(c.value or '').strip() for c in ws[1]]
    out=[]
    for row in ws.iter_rows(min_row=2,values_only=True):
        d=dict(zip(headers,row))
        if not d.get('Event_Code'): continue
        for col in ('Keyword','Exclude_Keyword'):
            d[col]=[w.strip() for w in re.split(r'[;；\n]+',str(d.get(col) or '')) if w.strip()]
        out.append(d)
    wb.close()
    return out

def term_hit(text: str,term: str)->bool:
    # Do not match short English substrings within longer words.
    if re.search(r'[A-Za-z]',term):
        return bool(re.search(r'(?<![A-Za-z])'+re.escape(term)+r'(?![A-Za-z])',text,re.I))
    return term.casefold() in text.casefold()

def candidate_rules(text: str,rules: list[dict]):
    out=[]
    cds=bool(CDS_PAT.search(text)); actual=bool(ACTUAL_DEFAULT.search(text))
    for r in rules:
        keywords=r.get('Keyword',[])
        if not keywords or str(r.get('Event_Code','')) in SYSTEM_CODES: continue
        if any(term_hit(text,x) for x in r.get('Exclude_Keyword',[])): continue
        hits=[w for w in keywords if term_hit(text,w)]
        if not hits: continue
        # CDS is not corporate default merely because 'default' is inside the instrument's name.
        default_only = hits and all(x.strip().casefold() in ('default','defaults','違約') for x in hits)
        code=str(r.get('Event_Code','')).upper()
        typename=str(r.get('Event_type','')).upper()
        if cds and not actual and default_only and ('DEFAULT' in code or '違約' in str(r.get('事件中文','')) or 'DEFAULT' in typename):
            continue
        out.append({'Event_Code':str(r.get('Event_Code','')),
                    '事件中文':str(r.get('事件中文','') or ''),
                    'Level':str(r.get('Level','') or ''),
                    'Action':str(r.get('Action','') or ''),
                    'hits':hits,
                    'priority':int(r.get('Primary_Priority') or 0)})
    out.sort(key=lambda r:-r['priority'])
    return out

def pick_rule(event_code: str,rules:list[dict]):
    for r in rules:
        if str(r.get('Event_Code','')).strip().casefold()==str(event_code).strip().casefold():
            return r
    return None
