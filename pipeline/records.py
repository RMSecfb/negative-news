"""Canonical news identity, non-destructive merges, manual decisions and shared Excel reports."""
from __future__ import annotations
import hashlib
import io
import json
import re
from datetime import datetime, timedelta, timezone
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment
from .common import clean_url, truth
from .repository import json_bytes, read_json

TPE = timezone(timedelta(hours=8))
FINAL_FIELDS = ['Ticker', 'Company', 'Event_Code', '事件中文', 'Level', 'Action', '覆核註記']
LEVEL_COLORS = {3: 'FFF9C4', 4: 'FFE0B2', 5: 'FFCDD2'}


def csv_frame(raw):
    if not raw:
        return pd.DataFrame()
    return pd.read_csv(io.BytesIO(raw), dtype=str, keep_default_na=False)


def published(value):
    from .market import publication_time
    value = str(value or '')
    dt = publication_time(value)
    if dt is None:
        try:
            # Original method one/two explicitly emit Taipei local time.
            dt = datetime.fromisoformat(value).replace(tzinfo=TPE)
        except ValueError:
            return None
    return dt.astimezone(TPE)


def normalize(frame):
    rows = []
    for row in frame.fillna('').to_dict('records'):
        r = {k: ('' if v is None else v) for k, v in row.items()}
        r['Market'] = str(r.get('Market') or 'US').upper()
        dt = published(r.get('Published Time', ''))
        if dt is None:
            continue
        r['Published Time'] = dt.isoformat(timespec='seconds')
        r['日期'] = dt.date().isoformat()
        r.setdefault('來源方法', '')
        r.setdefault('人工版本', '')
        r.setdefault('已人工覆核', 'False')
        if not r.get('最終判斷來源'):
            r['最終判斷來源'] = '人工覆核' if truth(r.get('已人工覆核')) else 'LLM' if r.get('LLM狀態') == 'OK' else '規則' if r.get('Event_Code') else '待補分析'
        material = clean_url(str(r.get('URL', ''))) or re.sub(r'\W+', '', str(r.get('Title', '')).casefold())
        r['article_id'] = hashlib.sha256((r['Market'] + '|' + material).encode()).hexdigest()[:24]
        issuer = str(r.get('Ticker') or r.get('Company') or '').casefold().strip()
        r['news_id'] = str(r.get('news_id') or hashlib.sha256((r['article_id'] + '|' + issuer).encode()).hexdigest()[:24])
        rows.append(r)
    return pd.DataFrame(rows)


def tags(*values):
    return '|'.join(sorted({x.strip() for v in values for x in re.split(r'[|;；]', str(v or '')) if x.strip()}))


def merge_news(old, incoming):
    """Only collapse same issuer+URL or exact headline; related articles remain separate."""
    records = normalize(old).to_dict('records')
    added = duplicates = 0
    for r in normalize(incoming).to_dict('records'):
        ticker = str(r.get('Ticker') or r.get('原始Ticker') or '').upper()
        company = str(r.get('Company', '')).casefold()
        candidates = []
        for i, prev in enumerate(records):
            if prev['Market'] != r['Market']:
                continue
            pt = str(prev.get('Ticker') or prev.get('原始Ticker') or '').upper()
            pc = str(prev.get('Company', '')).casefold()
            issuer_match = (ticker and pt and ticker == pt) or (company and pc and company == pc)
            same_article = prev['article_id'] == r['article_id']
            same_title = re.sub(r'\W+', '', str(prev.get('Title', '')).casefold()) == re.sub(r'\W+', '', str(r.get('Title', '')).casefold())
            if (same_article and (issuer_match or not ticker or not pt)) or (same_title and issuer_match and prev['日期'] == r['日期']):
                candidates.append(i)
        if len(candidates) != 1:
            records.append(r)
            added += 1
            continue
        index = candidates[0]
        prev = records[index]
        duplicates += 1
        combined = dict(prev)
        for field in ('Title', 'Summary', 'Source', 'FinBERT', 'Title_ZH', 'Rule_Event_Code', 'Rule_Level', 'Rule_Action', 'Rule_事件中文', '監控池'):
            if str(r.get(field, '')).strip():
                combined[field] = r[field]
        for field in ('Ticker', 'Company', 'Exchange', 'Ticker來源'):
            if not prev.get(field) and r.get(field):
                combined[field] = r[field]
        combined['來源方法'] = tags(prev.get('來源方法'), r.get('來源方法'))
        combined['監控池'] = tags(prev.get('監控池'), r.get('監控池'))
        # An analyzed incoming row may improve a pending row, but never a human decision.
        if not truth(prev.get('已人工覆核')) and r.get('LLM狀態') == 'OK' and prev.get('LLM狀態') != 'OK':
            protected = {'news_id', 'article_id', '來源方法', '監控池'}
            combined.update({k: v for k, v in r.items() if k not in protected})
        if truth(prev.get('已人工覆核')):
            for field in [*FINAL_FIELDS, '人工版本', '人工覆核人', '人工覆核時間', '已人工覆核', '分類狀態', '需要人工覆核', '最終判斷來源']:
                if field in prev:
                    combined[field] = prev[field]
        records[index] = combined
    return pd.DataFrame(records).fillna(''), {'新增': added, '重複合併': duplicates}


def major(frame):
    if frame.empty:
        return frame.copy()
    levels = pd.to_numeric(frame.get('Level', pd.Series('', index=frame.index)), errors='coerce')
    return frame[levels.between(3, 5)].copy()


def pending(frame):
    if frame.empty:
        return frame.copy()
    return frame[frame.get('分類狀態', pd.Series('', index=frame.index)).eq('待補分析')].copy()


def review(frame):
    if frame.empty:
        return frame.copy()
    return frame[frame.get('需要人工覆核', pd.Series('', index=frame.index)).map(truth)].copy()


def excel_bytes(frame, runs=None):
    """Application export: the same four report views for schedule, manual runs and queries."""
    out = io.BytesIO()
    preferred = ['日期', 'Market', 'Ticker', 'Company', 'Published Time', 'Title', '中文摘要', 'Level', '事件中文', 'Action', '最終判斷來源', '分類狀態', '已人工覆核', '來源方法', 'FinBERT', 'Rule_Event_Code', 'Rule_Level', 'LLM事件代碼', 'LLM辨識公司', 'Ticker來源', 'URL', 'Source']
    cols = [c for c in preferred if c in frame] + [c for c in frame if c not in preferred]
    frame = frame.reindex(columns=cols)
    sheets = {'完整新聞': frame, '主要負面事件': major(frame), '待人工覆核': review(frame), '待補分析': pending(frame), '執行紀錄': pd.DataFrame(runs or [])}
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        for name, data in sheets.items():
            safe = data.fillna('').copy()
            for col in safe.columns:
                safe[col] = safe[col].map(lambda v: "'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v)
            safe.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='1F4E78')
                cell.font = Font(color='FFFFFF', bold=True)
            level_col = list(data.columns).index('Level') + 1 if 'Level' in data else None
            for row in ws.iter_rows(min_row=2):
                color = None
                if level_col:
                    try:
                        color = LEVEL_COLORS.get(int(float(row[level_col - 1].value or 0)))
                    except (ValueError, TypeError):
                        pass
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if color:
                        cell.fill = PatternFill('solid', fgColor=color)
            from openpyxl.utils import get_column_letter
            for i, col in enumerate(data.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = 65 if col in ('Title', 'URL', 'Evidence') else 38 if col in ('Action', '中文摘要') else 22
    return out.getvalue()


def read_days(repository, start, end, market='ALL', revision=None, issuer_index=None):
    frames = []
    missing_days = []
    day = start
    while day <= end:
        raw = repository.read(f'news_daily/{day.isoformat()}.csv', revision)
        if raw is None:
            missing_days.append(day)
            raw = repository.read(f'ai_news_daily/{day.isoformat()}.csv', revision)
            legacy = repository.read(f'risk_news_daily/{day.isoformat()}.xlsx', revision)
            frame = normalize(csv_frame(raw))
            if legacy:
                old = pd.read_excel(io.BytesIO(legacy), dtype=str).fillna('')
                old['來源方法'] = '原有新聞流程'
                frame, _ = merge_news(frame, old)
        else:
            frame = normalize(csv_frame(raw))
        if not frame.empty:
            frames.append(frame)
        day += timedelta(days=1)
    # Legacy AI files are named by execution day, while stories may belong to
    # either preceding day. Read that bounded overlap when migrating old dates.
    missing_dates = {d.isoformat() for d in missing_days}
    for extra_day in sorted({d + timedelta(days=n) for d in missing_days for n in (1, 2)}):
        frame = normalize(csv_frame(repository.read(f'ai_news_daily/{extra_day.isoformat()}.csv', revision)))
        if not frame.empty:
            frames.append(frame[frame['日期'].isin(missing_dates)])
    result = pd.concat(frames, ignore_index=True).fillna('') if frames else pd.DataFrame()
    if not result.empty:
        result = result.drop_duplicates('news_id')
        result = result[result['日期'].between(start.isoformat(), end.isoformat())]
        if market != 'ALL':
            result = result[result['Market'] == market]
        result = result.sort_values('Published Time', ascending=False)
    # Migrate original manual overrides by URL/ticker without changing their files.
    overrides = csv_frame(repository.read('history/negative_news_manual_overrides.csv', revision)) if not result.empty else pd.DataFrame()
    if not overrides.empty and 'URL' in overrides:
        for override in overrides.to_dict('records'):
            mask = result['URL'].map(clean_url).eq(clean_url(override.get('URL', '')))
            if override.get('Ticker'):
                mask &= result['Ticker'].astype(str).str.upper().eq(str(override['Ticker']).upper())
            mask &= ~result.get('已人工覆核', pd.Series('', index=result.index)).map(truth)
            for field in ['事件中文', 'Level', 'Action']:
                if str(override.get(field, '')).strip():
                    result.loc[mask, field] = str(override[field])
            for field, value in {'已人工覆核': 'True', '分類狀態': '人工覆核', '需要人工覆核': 'False', '最終判斷來源': '人工覆核'}.items():
                result.loc[mask, field] = value
    if not result.empty:
        from pathlib import Path
        from .issuers import load_index
        issuer_index = issuer_index or load_index(Path(__file__).resolve().parents[1], repository)
        result = issuer_index.enrich(result)
    return result.reset_index(drop=True)


def manual_save(repository, original, edited, reviewer=''):
    """Compare only edited fields. A conflict never overwrites a newer human decision."""
    changes = []
    before = {str(r['news_id']): r for r in original.fillna('').to_dict('records')}
    for r in edited.fillna('').to_dict('records'):
        old = before.get(str(r.get('news_id', '')))
        if not old:
            continue
        if any(str(r.get(k, '')).strip() != str(old.get(k, '')).strip() for k in FINAL_FIELDS if k != 'Company'):
            try:
                level = float(r.get('Level', ''))
                if not level.is_integer() or not 1 <= level <= 5:
                    raise ValueError()
            except (ValueError, TypeError):
                raise ValueError('人工修改的 Level 必須是 1–5 的整數')
            changes.append((old, r))
    if not changes:
        return 0
    import uuid
    def build(read, revision):
        files, days = {}, {}
        for old, changed in changes:
            day = str(old['日期'])
            if day not in days:
                class Snapshot:
                    def read(self, path, revision=None):
                        return read(path)
                target_day = datetime.fromisoformat(day).date()
                days[day] = read_days(Snapshot(), target_day, target_day)
            df = days[day]
            matches = df.index[df['news_id'].astype(str) == str(old['news_id'])]
            if len(matches) != 1:
                raise ValueError('新聞識別已變更，請重新查詢後再編輯')
            index = matches[0]
            if str(df.loc[index].get('人工版本', '')) != str(old.get('人工版本', '')):
                raise ValueError('這筆新聞已有新的人工覆核，請重新載入後再修改')
            modified = [k for k in FINAL_FIELDS if k != 'Company' and str(changed.get(k,'')).strip() != str(old.get(k,'')).strip()]
            df.loc[index, '人工修改欄位'] = tags(df.loc[index].get('人工修改欄位',''), '|'.join(modified))
            for field in FINAL_FIELDS:
                if field != 'Company':
                    df.loc[index, field] = str(changed.get(field, '')).strip()
            from .issuers import symbol
            df.loc[index, 'Ticker'] = symbol(df.loc[index, 'Ticker'], str(df.loc[index, 'Market']))
            df.loc[index, 'Level'] = str(int(float(changed['Level'])))
            for field, value in {'已人工覆核': 'True', '需要人工覆核': 'False', '分類狀態': '人工覆核', '最終判斷來源': '人工覆核', '人工版本': uuid.uuid4().hex, '人工覆核人': reviewer, '人工覆核時間': datetime.now(TPE).isoformat()}.items():
                df.loc[index, field] = value
        for day, df in days.items():
            files[f'news_daily/{day}.csv'] = df.to_csv(index=False).encode('utf-8-sig')
            files[f'news_daily/{day}.xlsx'] = excel_bytes(df, read_json(read('news_runs/index.json'), {'runs': []})['runs'])
        return files, len(changes)
    return repository.transaction(build, 'Save manual news review without overwriting concurrent changes')
