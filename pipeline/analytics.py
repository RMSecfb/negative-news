"""Read-only historical news analytics and independently versioned SEC13F results."""
from __future__ import annotations
import io
import re
import pandas as pd
from openpyxl import load_workbook
from .records import excel_bytes

SEC_REPO = 'RMSecfb/negative-news'


def frequency_tables(frame):
    """Legacy spike thresholds; missing collection days are NOT assumed to be zero."""
    data = frame.copy()
    data['日期'] = pd.to_datetime(data['日期'], errors='coerce').dt.strftime('%Y-%m-%d')
    days = sorted(data['日期'].dropna().unique())
    levels = pd.to_numeric(data.get('Level', pd.Series(index=data.index, dtype=float)), errors='coerce')
    negative = data[levels.ge(1)].copy()
    keys = [c for c in ['Market', 'Ticker', '日期'] if c in negative]
    if 'news_id' in negative:
        negative = negative.drop_duplicates('news_id')
    counts = negative.groupby(keys, dropna=False).size().rename('則數').reset_index()
    anomalies = []
    if len(days) > 1:
        latest, baseline = days[-1], days[:-1]
        known = negative[negative['Ticker'].fillna('').ne('')]
        for (market, ticker), group in known.groupby(['Market', 'Ticker']):
            today = int(group['日期'].eq(latest).sum())
            avg = group['日期'].isin(baseline).sum() / len(baseline)
            if (avg == 0 and today >= 3) or (avg > 0 and today >= 2 and today >= 2 * avg):
                anomalies.append({'Market': market, 'Ticker': ticker, '比較日期': latest,
                                  '當日則數': today, '過去平均每日': round(avg, 3),
                                  '倍數': '新增' if avg == 0 else round(today / avg, 2)})
    return counts, pd.DataFrame(anomalies), days


def sec_paths(paths):
    return sorted([p for p in paths if re.fullmatch(
        r'institutional_13f/\d{4}Q[1-4]/13F_\d{4}Q[1-4]_Full_Universe\.xlsx', p)], reverse=True)


def parse_sec(data):
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype={'Ticker': str})
    required = {'Ticker Summary', '13F Detail', 'Run Log'}
    if not required.issubset(sheets):
        raise ValueError('13F 成果缺少 Ticker Summary、13F Detail 或 Run Log')
    summary = sheets['Ticker Summary']
    # The published 2026Q2 workbook has damaged Chinese headers. Only accept
    # its known 20-column layout after checking the stable English anchors.
    schema = ['Quarter', 'Quarter End', 'Ticker', 'Company', '追蹤Universe', '本季納入機構數',
              '本季持有機構數', '前季持有機構數', 'Current Shares', 'Previous Shares', '機構持股QoQ',
              '增持家數', '減持家數', '新建倉', '清倉', '持股不變家數', '淨增持家數', '持股量方向', '機構持股方向', '13F訊號']
    if any('\ufffd' in str(c) for c in summary.columns):
        anchors = {0:'Quarter', 1:'Quarter End', 2:'Ticker', 3:'Company', 8:'Current Shares', 9:'Previous Shares'}
        if len(summary.columns) != 20 or any(summary.columns[i] != name for i,name in anchors.items()):
            raise ValueError('13F 中文欄位受損且不符已知格式，需先修正來源檔')
        summary.columns = schema
        qoq = pd.to_numeric(summary['機構持股QoQ'], errors='coerce')
        summary['持股量方向'] = qoq.map(lambda n: '無前期比較資料' if pd.isna(n) else '持股增加' if n > 0 else '持股下降' if n < 0 else '持股不變')
        diff = pd.to_numeric(summary['增持家數'],errors='coerce') - pd.to_numeric(summary['減持家數'],errors='coerce')
        summary['機構持股方向'] = diff.map(lambda n: '增持家數較多' if n > 0 else '減持家數較多' if n < 0 else '增減持家數相同')
        summary['13F訊號'] = summary['持股量方向'] + '／' + summary['機構持股方向']
    if not {'Ticker', 'Quarter', 'Quarter End'}.issubset(summary):
        raise ValueError('13F 摘要缺少 Ticker、Quarter 或 Quarter End')
    summary['Ticker'] = summary['Ticker'].fillna('').str.strip().str.upper()
    if summary['Ticker'].eq('').any() or summary['Ticker'].duplicated().any():
        raise ValueError('13F 摘要含空白或重複 Ticker，停止串接以免放大新聞筆數')
    if summary['Quarter'].nunique() != 1 or summary['Quarter End'].nunique() != 1:
        raise ValueError('13F 摘要混合季度，停止串接')
    if not {'Ticker', 'Institution'}.issubset(sheets['13F Detail']):
        raise ValueError('13F 明細缺少 Ticker 或 Institution')
    return sheets


def join_sec(news, summary):
    data = news.copy()
    if data.empty:
        return data
    # Exact, validated US ticker only. Unknown/HK never matches US 13F by number/name.
    data['_sec_ticker'] = data['Ticker'].fillna('').astype(str).str.strip().str.upper().where(data['Market'].eq('US'), '')
    sec = summary.add_prefix('13F · ').rename(columns={'13F · Ticker': '_sec_ticker'})
    joined = data.merge(sec, on='_sec_ticker', how='left', validate='many_to_one').drop(columns='_sec_ticker')
    joined['13F · 對照狀態'] = joined['13F · Quarter'].notna().map({True: '有季度參考資料', False: '未涵蓋／代號待核對'})
    return joined


def combined_excel(news, sheets, provenance, runs=None):
    # Retain the existing five sheets and Level colors; append independent SEC data.
    wb = load_workbook(io.BytesIO(excel_bytes(news, runs or [])))
    frames = {'新聞與13F對照': join_sec(news, sheets['Ticker Summary']),
              '13F季度摘要': sheets['Ticker Summary'], '13F機構明細': sheets['13F Detail'],
              '13F執行紀錄': sheets['Run Log'], '資料來源與期間': pd.DataFrame([provenance])}
    for name, frame in frames.items():
        ws = wb.create_sheet(name)
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            values = []
            for value in row:
                if pd.isna(value):
                    value = None
                if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                    value = "'" + value
                values.append(value)
            ws.append(values)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
