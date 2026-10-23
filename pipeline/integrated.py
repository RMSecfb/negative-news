"""One merge/classification pipeline for scheduled and UI-requested runs."""
from __future__ import annotations
import hashlib
import json
import os
import re
import time
import uuid
from datetime import datetime, timedelta
from pathlib import Path
import pandas as pd
from .common import truth
from .events import load_rules, candidate_rules, pick_rule, CDS_PAT, ACTUAL_DEFAULT
from .issuers import IssuerIndex, load_companies, company_frame
from .llm import analyze, MODEL, PROMPT_VERSION
from .market import MARKET_QUERIES, priority
from .ingest import rss_fetch
from .records import TPE, normalize, merge_news, csv_frame, published, excel_bytes, read_days, pending, review, major
from .repository import read_json, json_bytes, GitHubRepository, LocalRepository

ROOT = Path(__file__).resolve().parents[1]
DECISION_FIELDS = ['Event_Code', '事件中文', 'Level', 'Action', '分類狀態', '需要人工覆核', '最終判斷來源']


def validate_request(request):
    r = dict(request)
    r.setdefault('id', uuid.uuid4().hex)
    if not re.fullmatch(r'[a-zA-Z0-9_-]{1,80}', str(r['id'])):
        raise ValueError('執行識別格式不正確')
    start, end = datetime.fromisoformat(r['start']), datetime.fromisoformat(r['end'])
    if start.tzinfo is None or end.tzinfo is None or not start < end:
        raise ValueError('開始／結束時間必須含時區，且開始時間早於結束時間')
    if end - start > timedelta(days=31):
        raise ValueError('單次人工擷取最多 31 天；查詢可跨更長期間')
    if r.get('market', 'US') not in ('US', 'HK') or r.get('scope', 'list') not in ('list', 'all'):
        raise ValueError('市場或標的範圍不正確')
    r.setdefault('market', 'US')
    r.setdefault('scope', 'list')
    r.setdefault('methods', ['方法一', '方法二'])
    r.setdefault('ai', True)
    r.setdefault('lists', ['SP500'] if r['market'] == 'US' else [])
    r.setdefault('companies', [])
    r.setdefault('created_at', datetime.now(TPE).isoformat())
    if not isinstance(r['ai'], bool) or not isinstance(r['methods'], list) or any(m not in ('方法一', '方法二') for m in r['methods']):
        raise ValueError('執行方法格式不正確')
    if not r['methods'] and not r['ai']:
        raise ValueError('請至少選擇一個新聞方法或 AI LLM')
    if len(r['companies']) > 3000:
        raise ValueError('自訂清單最多 3000 家公司')
    r['max_llm'] = min(100, max(0, int(r.get('max_llm', 35))))
    return r


def daily_request(now=None, max_llm=35):
    now = (now or datetime.now(TPE)).astimezone(TPE)
    end = now.replace(hour=8, minute=20, second=0, microsecond=0)
    if end > now:
        end -= timedelta(days=1)
    return validate_request({'id': 'daily-' + end.date().isoformat(), 'type': 'schedule',
                             'start': (end - timedelta(days=1)).isoformat(), 'end': end.isoformat(),
                             'market': 'US', 'scope': 'list', 'lists': ['SP500'],
                             'methods': ['方法一', '方法二'], 'ai': True,
                             'discover_market': True, 'max_llm': max_llm})


def event_discovery(start, end, market='US', pause=.3):
    import requests
    rows, errors = [], []
    session = requests.Session()
    after = (start.astimezone(TPE).date() - timedelta(days=1)).isoformat()
    before = (end.astimezone(TPE).date() + timedelta(days=1)).isoformat()
    for index, query in enumerate(MARKET_QUERIES):
        language = 'zh' if re.search(r'[\u4e00-\u9fff]', query) else 'en'
        if market == 'HK':
            query = query.replace('美股', '港股')
            if language == 'en':
                query += ' (Hong Kong OR HKEX)'
        query = f'({query}) after:{after} before:{before}'
        print(f'Market search {index + 1}/{len(MARKET_QUERIES)}', flush=True)
        try:
            for row in rss_fetch(query, language, session, 30):
                dt = published(row.get('Published Time', ''))
                if dt and start <= dt < end:
                    rows.append({**row, 'Market': market, '來源方法': '事件搜尋', '監控池': 'MARKET_DISCOVERY'})
        except Exception as exc:
            errors.append(f'事件搜尋 {index + 1}：{type(exc).__name__}')
        time.sleep(pause)
    return normalize(pd.DataFrame(rows)), errors


def reference_companies(repo, request):
    selected = load_companies(repo, request['lists'], request['companies'], request['market'])
    reference = load_companies(repo, ['SP500', 'DowJones30'])
    watch = repo / 'config' / 'extra_watchlist.csv'
    if watch.exists():
        reference = pd.concat([reference, company_frame(pd.read_csv(watch, dtype=str), 'US')], ignore_index=True)
    symbols = repo / 'config' / 'us_symbol_reference.csv'
    if symbols.exists():
        reference = pd.concat([reference, company_frame(pd.read_csv(symbols, dtype=str), 'US')], ignore_index=True)
    return selected, IssuerIndex(pd.concat([reference, selected], ignore_index=True))


def collect(repo, request, companies, progress=None):
    """Keep original US collectors/FinBERT. An all-market run uses the source pools directly."""
    import legacy_engine as engine
    engine.initialize()
    start, end = datetime.fromisoformat(request['start']), datetime.fromisoformat(request['end'])
    frames, errors, stages = [], [], []
    selected = request['scope'] == 'list'
    if selected and companies.empty:
        raise ValueError('請選擇或上傳至少一份有效公司列表')
    if '方法一' in request['methods']:
        print('Collect method one', flush=True)
        if selected and request['market'] == 'US':
            frame, failed, counts = engine.fetch_method1_news(companies, start, end)
        else:
            items, failed, counts = [], [], {}
            for source, fetcher in [('MoneyDJ', engine.method1_moneydj), ('經濟日報', engine.method1_udn), ('鉅亨網', engine.method1_cnyes), ('CNBC', engine.method1_cnbc)]:
                try:
                    for item in fetcher(start, end):
                        matches = engine.method2_matches(item['Title'], companies) if selected else [('', '')]
                        for ticker, company in matches:
                            items.append({'Ticker': ticker, 'Company': company, 'Title': item['Title'], 'Summary': item.get('Summary', ''), 'Published Time': item['Published'].isoformat(), 'URL': item['URL'], 'Source': item['Source']})
                except Exception as exc:
                    failed.append(f'{source}：{type(exc).__name__}')
            frame = pd.DataFrame(items)
            if request['market'] == 'HK':
                stages.append('港股方法一使用可用新聞池；原有美股 TradingView 逐公司補抓不套用港股')
        frame['Market'], frame['來源方法'] = request['market'], '方法一'
        frames.append(frame)
        errors.extend(failed)
        stages.append(f'方法一 {len(frame)} 筆')
    if '方法二' in request['methods'] and selected:
        print('Collect method two for every selected company', flush=True)
        if request['market'] == 'US':
            frame, failed, batches = engine.fetch_method2_raw(companies, start, end)
        else:
            import requests
            rows, failed = [], []
            session = requests.Session()
            for company in companies.to_dict('records'):
                query = f'"{company["Company"]}" after:{(start.date()-timedelta(days=1)).isoformat()} before:{(end.date()+timedelta(days=1)).isoformat()}'
                for lang in ('zh', 'en'):
                    try:
                        for item in rss_fetch(query, lang, session, 30):
                            dt = published(item['Published Time'])
                            if dt and start <= dt < end:
                                rows.append({**item, **{k: company[k] for k in ('Ticker', 'Company', 'Market')}})
                    except Exception as exc:
                        failed.append(f'Google News {company["Ticker"]}/{lang}：{type(exc).__name__}')
            frame = pd.DataFrame(rows)
        frame['Market'], frame['來源方法'] = request['market'], '方法二'
        frames.append(frame)
        errors.extend(failed)
        stages.append(f'方法二 {len(frame)} 筆；公司 {len(companies)} 家（非輪替）')
    discover = request.get('discover_market') or (not selected and '方法二' in request['methods']) or not request['methods']
    if discover:
        if selected and not request['methods']:
            # AI-only + company list: use original company RSS as the candidate source.
            frame, failed, _ = engine.fetch_method2_raw(companies, start, end)
            frame['Market'], frame['來源方法'] = request['market'], 'AI候選搜尋'
        else:
            frame, failed = event_discovery(start, end, request['market'])
            if '方法二' in request['methods'] and not selected:
                frame['來源方法'] = '方法二|事件搜尋'
        frames.append(frame)
        errors.extend(failed)
        stages.append(f'市場事件搜尋 {len(frame)} 筆')
    combined = pd.DataFrame()
    for frame in frames:
        combined, _ = merge_news(combined, frame)
    if not combined.empty and '方法二' in request['methods']:
        print(f'FinBERT on merged unique rows: {len(combined)}', flush=True)
        # The original engine translates Chinese titles before scoring and caches scores.
        scored, _ = engine.score_method2_finbert(combined)
        originals = combined.set_index('news_id')['Title'].to_dict()
        scored['FinBERT_Title'] = scored['Title']
        scored['Title'] = scored['news_id'].map(originals)
        combined = scored
        stages.append(f'FinBERT 完成 {len(combined)} 筆')
    return normalize(combined), errors, stages


def rule_fields(record, rules, core_sources=None):
    """Keep the legacy substring/priority rule result for audit and AI-off execution."""
    text = str(record.get('FinBERT_Title') or record.get('Title', '')).lower()
    hits = []
    for rule in rules:
        if any(str(x).lower() in text for x in rule.get('Exclude_Keyword', [])):
            continue
        if any(str(x).lower() in text for x in rule.get('Keyword', [])):
            hits.append(rule)
    hits.sort(key=lambda r: (-int(r.get('Primary_Priority') or 0), str(r.get('Event_ID') or '')))
    chosen = hits[0] if hits else None
    if chosen is None and core_sources is not None:
        source = str(record.get('Source', '')).strip().casefold()
        try:
            positive = float(record.get('FinBERT', '')) > 0
        except (TypeError, ValueError):
            positive = False
        fallback = 'NON_CORE_SOURCE' if source not in core_sources else 'NOT_RELEVANT' if positive else 'UNKNOWN'
        chosen = pick_rule(fallback, rules)
    return {f'Rule_{key}': str(chosen.get(key, '') or '') if chosen else '' for key in ('Event_Code', '事件中文', 'Level', 'Action')}


def fingerprint(record, rule_hash):
    fields = [MODEL, PROMPT_VERSION, rule_hash, record.get('Market', ''), record.get('Title', ''), record.get('Summary', '')]
    return hashlib.sha256(json.dumps(fields, ensure_ascii=False).encode()).hexdigest()


def apply_answer(record, answer, rules, index):
    r = dict(record)
    r.update({'LLM狀態': 'OK', 'API錯誤': '', 'LLM辨識公司': answer['company_mentioned'], 'LLM事件代碼': '|'.join(answer['event_codes']),
              '事件狀態': answer['event_status'], 'Confirmed_Default': str(answer['confirmed_default']), 'Evidence': answer['evidence'],
              '中文摘要': answer['summary_zh'], '最終判斷來源': 'LLM'})
    r = index.enrich(pd.DataFrame([r])).iloc[0].to_dict()
    chosen = next((pick_rule(c, rules) for c in answer['event_codes'] if pick_rule(c, rules)), None)
    for field in ('Event_Code', '事件中文', 'Level', 'Action'):
        r[field] = str(chosen.get(field, '') or '') if chosen and answer['relevant'] else ''
    unsupported_default = str(r.get('Event_Code', '')).upper() == 'DEFAULT' and (not answer['confirmed_default'] or answer['event_status'] != 'CONFIRMED')
    cds_conflict = bool(CDS_PAT.search(str(r.get('Title', ''))) and not ACTUAL_DEFAULT.search(str(r.get('Title', ''))) and answer['confirmed_default'])
    if unsupported_default or cds_conflict:
        r.update(Event_Code='', 事件中文='疑似違約／尚未確認（待覆核）', Level='', Action='人工核實新聞證據')
    try:
        major_event = float(r.get('Level') or 0) >= 4
    except ValueError:
        major_event = False
    evidence_valid = bool(answer['evidence']) and answer['evidence'] in (str(r.get('Title', '')) + '\n' + str(r.get('Summary', '')))
    needs = bool(answer['requires_review'] or not evidence_valid or unsupported_default or cds_conflict or
                 (answer['relevant'] and (not r.get('Ticker') or not r.get('Event_Code') or major_event or answer['confirmed_default'])))
    r['需要人工覆核'] = str(needs)
    r['分類狀態'] = '待人工覆核' if needs else ('已分類' if answer['relevant'] else '無關新聞')
    return r


def reserve(repository, run_id, requested, day, cap):
    path = f'news_usage/{day}.json'
    def build(read, revision):
        ledger = read_json(read(path), {'reservations': {}})
        entries = ledger['reservations']
        if run_id in entries:
            return {}, 0
        available = max(0, cap - sum(int(r['allocated']) for r in entries.values()))
        amount = min(requested, available)
        entries[run_id] = {'allocated': amount, 'used': 0}
        return {path: json_bytes(ledger)}, amount
    return repository.transaction(build, 'Reserve bounded shared LLM allowance')


def execute(repository, request, repo=ROOT, api_key='', collector=None, analyzer=None, interval=None):
    request = validate_request(request)
    run_id = request['id']
    start, end = datetime.fromisoformat(request['start']), datetime.fromisoformat(request['end'])
    selected, index = reference_companies(repo, request)
    rule_path = repo / 'Parameter_Event.xlsx'
    rules = load_rules(rule_path)
    from openpyxl import load_workbook
    wb = load_workbook(rule_path, read_only=True, data_only=True)
    sheet = wb['主要網站']
    headers = [str(c.value or '').strip() for c in sheet[1]]
    source_index = headers.index('Source')
    core_sources = {str(row[source_index]).strip().casefold() for row in sheet.iter_rows(min_row=2, values_only=True) if row[source_index]}
    wb.close()
    rule_hash = hashlib.sha256(rule_path.read_bytes()).hexdigest()[:16]
    usage_day = datetime.now(TPE).date().isoformat()
    report = {'id': run_id, 'type': request.get('type', 'manual'), 'status': 'running',
              'start': request['start'], 'end': request['end'], 'market': request['market'],
              'methods': '|'.join(request['methods']), 'ai': request['ai'], 'started_at': datetime.now(TPE).isoformat()}
    repository.transaction(lambda read, rev: ({f'news_runs/{run_id}.json': json_bytes(report)}, None), 'Start unified news run')
    print(f'Run {run_id}: {request["methods"]}, AI={request["ai"]}, {start} → {end}', flush=True)
    cache_path = f'news_cache/{usage_day[:7]}.json'
    cache = read_json(repository.read(cache_path))
    base = index.enrich(read_days(repository, start.astimezone(TPE).date(), (end - timedelta(microseconds=1)).astimezone(TPE).date()))
    if request.get('resume'):
        fresh, errors, stages = pd.DataFrame(), [], ['沿用已擷取資料；不重新抓取']
    else:
        fresh, errors, stages = (collector or collect)(repo, request, selected)
    fresh = index.enrich(normalize(fresh))
    combined, merge_stats = merge_news(base, fresh)
    if combined.empty:
        targets = pd.DataFrame()
    else:
        times = combined['Published Time'].map(published)
        targets = combined[(times >= start) & (times < end) & (combined['Market'] == request['market'])].copy()
        if request['scope'] == 'list' and not request.get('discover_market'):
            targets = targets[targets['Ticker'].isin(set(selected['Ticker']))]
    requested = request['max_llm'] if request['ai'] and api_key else 0
    allowance = reserve(repository, run_id, requested, usage_day, int(os.getenv('LLM_DAILY_MAX', '120'))) if requested else 0
    calls = reused = 0
    new_cache, decisions = {}, []
    limit_reached = False
    ordered = sorted(targets.to_dict('records'), key=priority, reverse=True)
    pause = float(os.getenv('GROQ_MIN_INTERVAL_SECONDS', '20')) if interval is None else interval
    for record in ordered:
        r = {**record, **rule_fields(record, rules, core_sources)}
        if truth(r.get('已人工覆核')):
            r.update(最終判斷來源='人工覆核', 分類狀態='人工覆核', 需要人工覆核='False')
            decisions.append(r)
            continue
        key = fingerprint(r, rule_hash)
        if not request['ai']:
            for field in ('Event_Code', '事件中文', 'Level', 'Action'):
                r[field] = r.get('Rule_' + field, '')
            system = r.get('Event_Code') in ('NOT_RELEVANT', 'NON_CORE_SOURCE')
            unknown = not r.get('Event_Code') or r.get('Event_Code') == 'UNKNOWN'
            r.update(最終判斷來源='規則', 分類狀態='無關新聞' if system else '待人工覆核' if unknown else '已分類', 需要人工覆核=str(unknown), LLM狀態='DISABLED')
        elif key in cache:
            r = apply_answer(r, cache[key]['answer'], rules, index)
            r['分析快取'] = '沿用'
            reused += 1
        elif r.get('LLM狀態') == 'OK' and r.get('Prompt版本') == PROMPT_VERSION and str(r.get('LLM模型', '')) == MODEL and (not r.get('分析指紋') or r['分析指紋'] == key):
            # Valid old daily results remain available when first upgrading to the unified store.
            answer = {'company_mentioned': r.get('LLM辨識公司', ''), 'event_codes': str(r.get('LLM事件代碼', '')).split('|') if r.get('LLM事件代碼') else [],
                      'event_status': r.get('事件狀態') or 'UNCLEAR', 'confirmed_default': truth(r.get('Confirmed_Default')),
                      'evidence': r.get('Evidence', ''), 'summary_zh': r.get('中文摘要', ''),
                      'relevant': r.get('分類狀態') != '無關新聞', 'requires_review': truth(r.get('需要人工覆核'))}
            r = apply_answer(r, answer, rules, index)
            new_cache[key] = {'answer': answer}
            reused += 1
        elif not api_key or calls >= allowance or limit_reached:
            r.update(LLM狀態='NO_API_KEY' if not api_key else 'QUOTA_PENDING', 分類狀態='待補分析', 需要人工覆核='False', 最終判斷來源='待補分析')
            for field in ('Event_Code', '事件中文', 'Level', 'Action'):
                r[field] = ''
        else:
            if calls and pause:
                time.sleep(pause)
            codes = [str(rule['Event_Code']) for rule in rules if str(rule['Event_Code']) not in ('UNKNOWN', 'NOT_RELEVANT', 'NON_CORE_SOURCE')]
            labels = {str(rule['Event_Code']): str(rule.get('事件中文') or '') for rule in rules}
            answer, usage, error = (analyzer or analyze)(r, codes, api_key, labels)
            calls += 1
            print(f'LLM {calls}/{allowance}: {r.get("Ticker", "")} | {r.get("Title", "")[:75]}', flush=True)
            if answer is None:
                limit_reached = bool(re.search(r'429|rate.?limit', str(error), re.I))
                r.update(LLM狀態='API_ERROR', API錯誤=str(error)[:240], 分類狀態='待補分析', 最終判斷來源='待補分析', 需要人工覆核='False')
                for field in ('Event_Code', '事件中文', 'Level', 'Action'):
                    r[field] = ''
            else:
                r = apply_answer(r, answer, rules, index)
                r.update(input_tokens=usage.get('input_tokens', 0), output_tokens=usage.get('output_tokens', 0))
                new_cache[key] = {'answer': answer}
        r.update(分析指紋=key, LLM模型=MODEL if request['ai'] else '', Prompt版本=PROMPT_VERSION if request['ai'] else '', 最近執行=run_id)
        decisions.append(r)
    analyzed = pd.DataFrame(decisions).fillna('')
    report.update(status='success', finished_at=datetime.now(TPE).isoformat(), llm_calls=calls, cache_reused=reused,
                  pending=len(pending(analyzed)), review=len(review(analyzed)), major=len(major(analyzed)),
                  selected_company_count=len(selected), collected=len(fresh), rows=len(targets), errors=errors,
                  stages=stages, quota_reserved=allowance, **merge_stats)
    # Re-read in the final transaction so a manual review made during crawling always wins.
    def publish(read, revision):
        files = {}
        run_index = read_json(read('news_runs/index.json'), {'runs': []})
        run_index['runs'] = sorted([r for r in run_index['runs'] if r.get('id') != run_id] + [report], key=lambda r: r.get('started_at', ''), reverse=True)[:100]
        if not combined.empty:
            replacements = {str(r['news_id']): r for r in analyzed.to_dict('records')}
            persisted = pd.DataFrame([replacements.get(str(r['news_id']), r) for r in combined.to_dict('records')]).fillna('')
            for day, group in persisted.groupby('日期'):
                latest = normalize(csv_frame(read(f'news_daily/{day}.csv') or read(f'ai_news_daily/{day}.csv')))
                if not latest.empty:
                    latest = latest[latest['日期'] == day]
                merged, _ = merge_news(latest, group)
                updates = {str(r['news_id']): r for r in group.to_dict('records')}
                result = []
                for row in merged.to_dict('records'):
                    update = updates.get(str(row['news_id']))
                    if update and not truth(row.get('已人工覆核')):
                        row.update(update)
                    result.append(row)
                frame = pd.DataFrame(result).fillna('')
                files[f'news_daily/{day}.csv'] = frame.to_csv(index=False).encode('utf-8-sig')
                files[f'news_daily/{day}.xlsx'] = excel_bytes(frame, run_index['runs'])
        latest_cache = read_json(read(cache_path))
        latest_cache.update(new_cache)
        files[cache_path] = json_bytes(latest_cache)
        ledger_path = f'news_usage/{usage_day}.json'
        ledger = read_json(read(ledger_path), {'reservations': {}})
        if run_id in ledger['reservations']:
            ledger['reservations'][run_id] = {'allocated': calls, 'used': calls}
            files[ledger_path] = json_bytes(ledger)
        files[f'news_runs/{run_id}.json'] = json_bytes(report)
        files['news_runs/index.json'] = json_bytes(run_index)
        return files, report
    return repository.transaction(publish, f'Merge unified news results: {run_id}')


def enqueue(repository, request):
    request = validate_request(request)
    path = f'news_requests/{request["id"]}.json'
    repository.transaction(lambda read, revision: ({path: json_bytes(request)}, request['id']), 'Queue manual unified news run')
    return request['id']


def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--event', default=os.getenv('GITHUB_EVENT_NAME', 'workflow_dispatch'))
    parser.add_argument('--max-llm', type=int, default=35)
    parser.add_argument('--resume', action='store_true')
    args = parser.parse_args()
    repository = GitHubRepository(os.environ['GITHUB_REPOSITORY'], os.environ['GITHUB_TOKEN'])
    requests = []
    if args.event in ('schedule', 'workflow_dispatch') and not args.resume:
        request = daily_request(max_llm=args.max_llm)
        if args.event == 'workflow_dispatch':
            request['id'] = 'manual-schedule-' + uuid.uuid4().hex
        requests.append(request)
    if args.resume:
        now = datetime.now(TPE)
        requests.append({'id': 'resume-' + uuid.uuid4().hex, 'start': (now-timedelta(days=2)).isoformat(), 'end': now.isoformat(), 'market': 'US', 'scope': 'all', 'lists': [], 'methods': [], 'ai': True, 'resume': True, 'max_llm': args.max_llm})
    # Every trigger drains pending UI requests. Requests already completed are idempotent.
    pending_requests = []
    for path in repository.list('news_requests/'):
        if not path.endswith('.json'):
            continue
        request = read_json(repository.read(path))
        if not repository.read(f'news_runs/{request.get("id", "")}.json'):
            pending_requests.append(request)
    requests.extend(sorted(pending_requests, key=lambda r: r.get('created_at', '')))
    failures = []
    for request in requests:
        completed = read_json(repository.read(f'news_runs/{request["id"]}.json'))
        if completed.get('status') == 'success':
            continue
        try:
            result = execute(repository, request, api_key=os.getenv('GROQ_API_KEY', ''))
            print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
        except Exception as exc:
            failed = {'id': request['id'], 'status': 'failed', 'error': str(exc)[:400], 'finished_at': datetime.now(TPE).isoformat()}
            repository.transaction(lambda read, rev: ({f'news_runs/{request["id"]}.json': json_bytes(failed)}, None), 'Record failed unified news run')
            failures.append(request['id'])
            print(f'ERROR {request["id"]}: {type(exc).__name__}: {exc}', flush=True)
    if failures:
        raise SystemExit('Failed requests: ' + ','.join(failures))


if __name__ == '__main__':
    main()
