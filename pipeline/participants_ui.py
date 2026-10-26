"""Institutional/insider view over the results produced by sec_filings/."""
import io
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from .participants import read_snapshot, annotate, parse_form4, SOURCE_REPO
from .analytics import parse_sec, combined_excel
from .records import excel_bytes


def export_participants(news, snapshot, workbooks, runs=None):
    manifest, summaries = snapshot
    news = annotate(news,snapshot)
    provenance = {'來源Repository':SOURCE_REPO,'同步時間':manifest.get('synced_at',''),
                  '來源版本':manifest.get('checked_source_sha',''),
                  '對照用途':'所選成果期間參考，非歷史當時資訊回測'}
    raw = (combined_excel(news,workbooks['13f']['sheets'],provenance,runs)
           if '13f' in workbooks else excel_bytes(news,runs))
    wb = load_workbook(io.BytesIO(raw))
    extra = {'機構內部人資料狀態':pd.DataFrame([{'類別':kind,**meta} for kind,meta in manifest.get('datasets',{}).items()])}
    if 'form4' in workbooks:
        for name, frame in workbooks['form4']['sheets'].items():
            extra['Form4 '+name] = frame
    for name,frame in extra.items():
        ws = wb.create_sheet(name[:31])
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False,name=None):
            values = []
            for value in row:
                if pd.isna(value):
                    value = None
                if isinstance(value,str) and value.startswith(('=','+','-','@')):
                    value = "'"+value
                values.append(value)
            ws.append(values)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render(repository, news):
    import streamlit as st
    st.subheader('機構內部人交易')
    st.caption('13F 機構季度持股 ＋ Form 4 內部人交易；由本專案 sec_filings 自動向 SEC 擷取，不改變新聞 Level。')
    st.link_button('開啟 13F／Form 4 擷取程式與原始成果', 'https://github.com/'+SOURCE_REPO+'/tree/main/sec_filings')
    st.caption('擷取排程（台北時間）：Form 4 週二至週六 06:00、13F 週一至週五 06:30，擷取後立即同步到這裡；07:00 另有保底同步。13F 為季度資料，新季度要多數機構申報後才會更新。')
    refresh = st.button('重新讀取機構內部人成果',key='refresh_participants')
    if refresh or 'participants_snapshot' not in st.session_state:
        try:
            st.session_state['participants_snapshot'] = read_snapshot(repository)
        except Exception as exc:
            st.error(f'機構內部人資料暫時無法讀取：{exc}')
            return
    snapshot = st.session_state['participants_snapshot']
    manifest,summaries = snapshot
    st.caption('上次同步：'+manifest.get('synced_at','尚未同步')+'；新聞表格以同一次查詢取得的摘要為準，重新查詢可更新。')
    st.warning('13F 顯示季度；Form 4 顯示分析起迄與資料品質。這是最新已匯入成果的參考對照，並非新聞發生當時已公開資訊回測；未涵蓋不等於零持股或沒有交易。')
    workbooks = {}
    for kind,label in [('13f','13F 機構持股'),('form4','Form 4 內部人交易')]:
        with st.expander(label,expanded=True):
            meta = manifest.get('datasets',{}).get(kind,{})
            if kind not in summaries:
                st.info(('尚無已發布的 Form 4 成果。' if kind=='form4' else '尚無已同步的13F成果。')+' '+meta.get('error',''))
                continue
            st.caption(f'來源：{meta.get("source_path", "")} · 版本 {meta.get("source_sha", "")[:12]}')
            if meta.get('source_status') != 'available':
                st.warning('本次同步未取得有效新成果，以下保留上次成功版本。'+meta.get('error',''))
            summary = summaries[kind]
            st.dataframe(summary,hide_index=True,use_container_width=True)
            path = meta.get('workbook_path','')
            cache_key = kind+':'+str(meta.get('source_sha'))+':'+path
            cached = st.session_state.setdefault('participant_workbooks',{})
            if path and cache_key not in cached:
                try:
                    if not path.startswith('institutional_insider/') or '..' in path:
                        raise ValueError('成果路徑無效')
                    raw = repository.read(path,manifest.get('_snapshot_revision'))
                    sheets = parse_sec(raw) if kind=='13f' else parse_form4(raw)
                    cached[cache_key] = {'raw':raw,'sheets':sheets}
                except Exception as exc:
                    st.error(f'明細檔暫時無法讀取：{exc}')
            if cache_key in cached:
                book = cached[cache_key]
                workbooks[kind] = book
                st.download_button('下載 '+label+' 原始 Excel',book['raw'],Path(path).name,key='download_participant_'+kind)
                detail_name = '13F Detail' if kind=='13f' else 'Form4 Detail'
                with st.expander(label+'：明細與執行品質'):
                    st.dataframe(book['sheets'][detail_name],hide_index=True)
                    st.dataframe(book['sheets']['Run Log'],hide_index=True)
                    if kind=='form4':
                        st.caption('NPR=(P筆數−S筆數)/(P筆數＋S筆數)。非P/S交易不納入；10b5-1保留標記，Form4/A及主體未辨識會提示覆核。')
                        for name in ['Excluded Transactions','Methodology']:
                            if name in book['sheets']:
                                st.dataframe(book['sheets'][name],hide_index=True)
    if not news.empty:
        st.markdown('#### 新聞與機構內部人交易對照')
        compared = annotate(news,snapshot)
        cols = [c for c in ['日期','Market','Ticker','Title','Level','13F 機構動向','Form4 內部人動向'] if c in compared]
        st.dataframe(compared[cols],hide_index=True,use_container_width=True)
        st.download_button('下載新聞＋機構內部人交易整合 Excel',export_participants(news,snapshot,workbooks,st.session_state.get('news_query_runs',[])), 'Negative_News_Institutional_Insider.xlsx',key='participants_export')
    else:
        st.info('先在查詢新聞期間取得新聞，這裡即可交叉查看並下載整合 Excel。')
